# 把 SOP Markdown 转成可打印/分发的 Word（含表格、标题层级、列表）
import re
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Cm

SRC = 'docs/SOP-结算平台试行.md'
OUT = 'docs/SOP-结算平台试行.docx'

doc = Document()
# 基础样式：中文友好
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(
    __import__('docx').oxml.ns.qn('w:eastAsia'), '微软雅黑')
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2.2)
    s.top_margin = s.bottom_margin = Cm(2.0)


def clean(text):
    # 去掉 markdown 强调标记，保留文字
    text = text.replace('**', '').replace('`', '')
    text = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text)
    return text.strip()


def add_table(rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = 'Light Grid Accent 1'
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(cols):
            txt = clean(row[j]) if j < len(row) else ''
            p = cells[j].paragraphs[0]
            p.text = txt
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.font.bold = True
    doc.add_paragraph()


lines = open(SRC, encoding='utf-8').read().split('\n')
i = 0
table_buf = []

while i < len(lines):
    line = lines[i].rstrip()

    # 表格聚合
    if line.startswith('|'):
        cells = [c.strip() for c in line.strip('|').split('|')]
        if not all(re.fullmatch(r':?-{2,}:?', c or '---') for c in cells):
            table_buf.append(cells)
        i += 1
        continue
    if table_buf:
        add_table(table_buf)
        table_buf = []

    if not line.strip():
        i += 1
        continue

    if line.startswith('#'):
        level = len(line) - len(line.lstrip('#'))
        text = clean(line.lstrip('#').strip())
        if level == 1:
            h = doc.add_heading(text, level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            doc.add_heading(text, level=min(level, 4))
        i += 1
        continue

    if line.startswith('>'):
        p = doc.add_paragraph(clean(line.lstrip('>')))
        p.paragraph_format.left_indent = Cm(0.6)
        for r in p.runs:
            r.font.color.rgb = RGBColor(0x60, 0x60, 0x60)
            r.font.size = Pt(9.5)
        i += 1
        continue

    if re.match(r'^\s*[-*]\s+', line):
        p = doc.add_paragraph(clean(re.sub(r'^\s*[-*]\s+', '', line)), style='List Bullet')
        i += 1
        continue

    if re.match(r'^\s*\d+\.\s+', line):
        p = doc.add_paragraph(clean(re.sub(r'^\s*\d+\.\s+', '', line)), style='List Number')
        i += 1
        continue

    if line.strip() == '---':
        i += 1
        continue

    doc.add_paragraph(clean(line))
    i += 1

if table_buf:
    add_table(table_buf)

doc.save(OUT)
print('已生成', OUT)
